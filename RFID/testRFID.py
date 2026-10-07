import socket
import struct
import threading
import time
import tkinter as tk
from tkinter import scrolledtext, messagebox
from datetime import datetime
from openpyxl import Workbook, load_workbook
import os

EXCEL_FILE = "rfid_database.xlsx"

BG     = "#070b10"
PANEL  = "#0d1218"
HEADER = "#101820"
TEXT   = "#00ff9c"
ACCENT = "#00d4ff"
WARN   = "#ffcc00"
ERROR  = "#ff5555"
DIM    = "#3a4a55"

class RFIDServer:
    def __init__(self, root):
        self.root = root
        self.root.title("RFID Gate Control System  •  TCP")
        self.root.geometry("1050x720")
        self.root.configure(bg=BG)

        self.server_socket = None
        self.running = False
        self.entry_times = {}
        self.connected_devices = {}

        self.setup_excel()
        self.build_ui()
        self.root.after(1000, self.update_devices_display)

    def setup_excel(self):
        if not os.path.exists(EXCEL_FILE):
            wb = Workbook()
            ws = wb.active
            ws.title = "ValidIDs"
            ws.append(["RFID", "Name"])
            ws.append(["F1C4F105", "Test User 1"])
            ws.append(["F0E7125F", "Test User 2"])
            ws2 = wb.create_sheet("Log")
            ws2.append(["Timestamp", "Gate", "RFID", "Name", "Action", "Duration (min)"])
            wb.save(EXCEL_FILE)
        self.wb = load_workbook(EXCEL_FILE)
        self.ws_valid = self.wb["ValidIDs"]
        self.ws_log = self.wb["Log"]

    def is_valid_id(self, rfid_hex):
        for row in self.ws_valid.iter_rows(min_row=2, values_only=True):
            if row[0] and str(row[0]).upper() == rfid_hex.upper():
                return True, row[1] if len(row) > 1 else "Unknown"
        return False, None

    def log_event(self, gate, rfid_hex, name, action, duration=None):
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.ws_log.append([now, gate, rfid_hex, name, action, duration])
        self.wb.save(EXCEL_FILE)

    def make_reply(self, status: int, rfid: int) -> bytes:
        packet = bytearray(9)
        packet[0] = (status >> 24) & 0xFF
        packet[1] = (status >> 16) & 0xFF
        packet[2] = (status >>  8) & 0xFF
        packet[3] = (status      ) & 0xFF
        packet[4:8] = struct.pack(">I", rfid)
        packet[8] = sum(packet[0:8]) & 0xFF
        return bytes(packet)

    def parse_packet(self, data: bytes):
        if len(data) != 9:
            return None
        if (sum(data[0:8]) & 0xFF) != data[8]:
            return None
        gate = struct.unpack(">I", data[0:4])[0]
        rfid = struct.unpack(">I", data[4:8])[0]
        return gate, rfid

    def handle_client(self, client_sock, addr):
        try:
            data = client_sock.recv(32)
            if not data:
                return

            parsed = self.parse_packet(data)
            if not parsed:
                self.root.after(0, self.add_log, f"Invalid packet from {addr[0]}", WARN)
                return

            gate, rfid = parsed
            rfid_hex = f"{rfid:08X}"

            # Update online status
            self.connected_devices[gate] = time.time()

            if rfid == 0:
                # Heartbeat
                self.root.after(0, self.add_log, f"Heartbeat  Gate {gate}  ({addr[0]})", ACCENT)
                client_sock.sendall(self.make_reply(1, 0))
                return

            # Real card
            valid, name = self.is_valid_id(rfid_hex)
            status = 1 if valid else 2
            client_sock.sendall(self.make_reply(status, rfid))
            self.root.after(0, self.add_log, f"Reply Status={status} → Gate {gate}", ACCENT)

            now = datetime.now()
            if gate == 1:
                self.entry_times[rfid_hex] = now
                self.log_event(gate, rfid_hex, name or "---", "ENTRY")
                msg = f"ENTRY   Gate {gate}   ID {rfid_hex}   {name or 'NOT FOUND'}"
            elif gate == 2:
                duration_min = None
                duration_str = "N/A"
                if rfid_hex in self.entry_times:
                    delta = now - self.entry_times[rfid_hex]
                    duration_min = round(delta.total_seconds() / 60, 1)
                    duration_str = f"{duration_min} min"
                    del self.entry_times[rfid_hex]
                self.log_event(gate, rfid_hex, name or "---", "EXIT", duration_min)
                msg = f"EXIT    Gate {gate}   ID {rfid_hex}   {name or 'NOT FOUND'}   Time: {duration_str}"
            else:
                msg = f"Gate {gate}   ID {rfid_hex}   {name or 'NOT FOUND'}"

            color = TEXT if valid else ERROR
            self.root.after(0, self.add_log, msg, color)

        except Exception as e:
            self.root.after(0, self.add_log, f"Client error: {e}", ERROR)
        finally:
            try:
                client_sock.close()
            except:
                pass

    def server_loop(self):
        while self.running:
            try:
                client_sock, addr = self.server_socket.accept()
                threading.Thread(target=self.handle_client, args=(client_sock, addr), daemon=True).start()
            except OSError:
                break
            except Exception as e:
                if self.running:
                    self.root.after(0, self.add_log, f"Server error: {e}", ERROR)

    def update_devices_display(self):
        now = time.time()
        active = []
        for gate, last in list(self.connected_devices.items()):
            if now - last < 12:
                active.append(f"Gate {gate:2d}   •   Online   •   Last seen {int(now - last)}s ago")
            else:
                del self.connected_devices[gate]

        self.devices_text.config(state="normal")
        self.devices_text.delete("1.0", tk.END)
        if active:
            for line in sorted(active):
                self.devices_text.insert(tk.END, line + "\n")
        else:
            self.devices_text.insert(tk.END, "No devices connected\n")
        self.devices_text.config(state="disabled")
        self.root.after(1000, self.update_devices_display)

    def connect(self):
        if self.running:
            return
        ip = self.ip_entry.get().strip()
        try:
            port = int(self.port_entry.get().strip())
        except:
            messagebox.showerror("Error", "Port must be a number")
            return

        try:
            self.server_socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            self.server_socket.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            self.server_socket.bind((ip, port))
            self.server_socket.listen(5)
        except Exception as e:
            messagebox.showerror("Bind Error", str(e))
            return

        self.running = True
        threading.Thread(target=self.server_loop, daemon=True).start()

        self.status_var.set(f"CONNECTED  •  TCP {ip}:{port}")
        self.status_label.config(fg=TEXT)
        self.btn_connect.config(state="disabled")
        self.btn_disconnect.config(state="normal")
        self.ip_entry.config(state="disabled")
        self.port_entry.config(state="disabled")
        self.add_log(f"TCP Server started on {ip}:{port}", ACCENT)

    def disconnect(self):
        self.running = False
        if self.server_socket:
            try:
                self.server_socket.close()
            except:
                pass
            self.server_socket = None
        self.connected_devices.clear()

        self.status_var.set("DISCONNECTED")
        self.status_label.config(fg=ERROR)
        self.btn_connect.config(state="normal")
        self.btn_disconnect.config(state="disabled")
        self.ip_entry.config(state="normal")
        self.port_entry.config(state="normal")
        self.add_log("TCP Server stopped", WARN)

    def build_ui(self):
        header = tk.Frame(self.root, bg=HEADER, height=52)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header, text="RFID GATE CONTROL SYSTEM  •  TCP", font=("Consolas", 17, "bold"),
                 bg=HEADER, fg=ACCENT).pack(side="left", padx=18, pady=11)

        control = tk.Frame(self.root, bg=PANEL, height=68)
        control.pack(fill="x", padx=12, pady=(10, 4))
        control.pack_propagate(False)

        tk.Label(control, text="LISTEN IP", font=("Consolas", 9), bg=PANEL, fg=DIM).place(x=18, y=6)
        self.ip_entry = tk.Entry(control, font=("Consolas", 12), bg="#0a1016", fg=TEXT,
                                 insertbackground=TEXT, relief="flat", width=15)
        self.ip_entry.insert(0, "0.0.0.0")
        self.ip_entry.place(x=18, y=26, height=28)

        tk.Label(control, text="PORT", font=("Consolas", 9), bg=PANEL, fg=DIM).place(x=190, y=6)
        self.port_entry = tk.Entry(control, font=("Consolas", 12), bg="#0a1016", fg=TEXT,
                                   insertbackground=TEXT, relief="flat", width=7)
        self.port_entry.insert(0, "5005")
        self.port_entry.place(x=190, y=26, height=28)

        self.btn_connect = tk.Button(control, text="CONNECT", font=("Consolas", 11, "bold"),
                                     bg="#0a3a2a", fg=TEXT, relief="flat", command=self.connect)
        self.btn_connect.place(x=300, y=20, width=105, height=34)

        self.btn_disconnect = tk.Button(control, text="DISCONNECT", font=("Consolas", 11, "bold"),
                                        bg="#3a1a1a", fg=ERROR, relief="flat", state="disabled",
                                        command=self.disconnect)
        self.btn_disconnect.place(x=420, y=20, width=115, height=34)

        self.status_var = tk.StringVar(value="DISCONNECTED")
        self.status_label = tk.Label(control, textvariable=self.status_var,
                                     font=("Consolas", 12, "bold"), bg=PANEL, fg=ERROR)
        self.status_label.place(x=560, y=26)

        main = tk.Frame(self.root, bg=BG)
        main.pack(fill="both", expand=True, padx=12, pady=4)

        left = tk.Frame(main, bg=PANEL)
        left.pack(side="left", fill="both", expand=True)

        tk.Label(left, text="LIVE EVENT LOG", font=("Consolas", 9), bg=PANEL, fg=DIM).pack(anchor="w", padx=8, pady=(6,0))
        self.log = scrolledtext.ScrolledText(left, font=("Consolas", 11), bg="#080c11",
                                             fg=TEXT, relief="flat", height=22)
        self.log.pack(fill="both", expand=True, padx=6, pady=6)
        self.log.tag_config("ok", foreground=TEXT)
        self.log.tag_config("err", foreground=ERROR)
        self.log.tag_config("info", foreground=ACCENT)
        self.log.tag_config("warn", foreground=WARN)

        right = tk.Frame(main, bg=PANEL, width=280)
        right.pack(side="right", fill="y", padx=(8,0))
        right.pack_propagate(False)

        tk.Label(right, text="CONNECTED DEVICES", font=("Consolas", 9), bg=PANEL, fg=DIM).pack(anchor="w", padx=8, pady=(6,0))
        self.devices_text = scrolledtext.ScrolledText(right, font=("Consolas", 11), bg="#080c11",
                                                      fg=TEXT, relief="flat", height=22, state="disabled")
        self.devices_text.pack(fill="both", expand=True, padx=6, pady=6)

        bottom = tk.Frame(self.root, bg=HEADER, height=28)
        bottom.pack(fill="x")
        tk.Label(bottom, text="Gate 1 = ENTRY   |   Gate 2 = EXIT   |   Heartbeat every 5s   |   Online timeout 12s",
                 font=("Consolas", 9), bg=HEADER, fg=DIM).pack(side="left", padx=14, pady=4)

    def add_log(self, message, color=TEXT):
        timestamp = datetime.now().strftime("%H:%M:%S")
        tag = "err" if color == ERROR else "info" if color == ACCENT else "warn" if color == WARN else "ok"
        self.log.insert("end", f"[{timestamp}]  {message}\n", tag)
        self.log.see("end")

    def on_close(self):
        self.disconnect()
        self.root.destroy()

if __name__ == "__main__":
    root = tk.Tk()
    app = RFIDServer(root)
    root.protocol("WM_DELETE_WINDOW", app.on_close)
    root.mainloop()